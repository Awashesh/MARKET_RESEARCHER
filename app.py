import os, io
from flask import Flask, render_template, request, jsonify, send_file
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import core
from kotak_live import KotakLiveManager

app=Flask(__name__); live=KotakLiveManager()
def ok(data=None,**kw):
    p={'ok':True};
    if data is not None:p['data']=data
    p.update(kw); return jsonify(p)
def fail(e,code=400): return jsonify({'ok':False,'error':str(e)}),code
def rec(df): return df.replace({float('nan'):None}).where(pd.notnull(df),None).to_dict('records')
def _int(name,default,lo,hi):
    try: v=int((request.args.get(name) or '').strip() or default)
    except (TypeError,ValueError): v=default
    return max(lo,min(v,hi))
def _float(name):
    x=(request.args.get(name) or '').strip()
    if not x:return None
    try:return float(x)
    except ValueError:return None

@app.get('/')
def index(): return render_template('index.html')
@app.get('/api/health')
def health(): return ok({'app':'Kotak Market Studio V2','kotak':live.status()})
@app.get('/api/kotak/status')
def status(): return ok(live.status())
@app.post('/api/kotak/login')
def login():
    try:
        t=str((request.get_json(silent=True) or {}).get('totp','')).strip()
        if len(t)!=6 or not t.isdigit(): return fail('Current 6-digit TOTP enter karein.')
        live.login_and_start(t); return ok(live.status())
    except Exception as e:return fail(e)
@app.get('/api/live')
def live_data(): return ok(live.snapshot())
@app.get('/api/quotes')
def quotes():
    try:
        s=[x.strip().upper() for x in request.args.get('symbols','RELIANCE,TCS,HDFCBANK,ICICIBANK,INFY').split(',') if x.strip()]
        return ok(live.rest_quotes(s[:40]))
    except Exception as e:return fail(e)
@app.post('/api/live/watch')
def watch():
    try:
        s=(request.get_json(silent=True) or {}).get('symbols',[]); live.set_watchlist(s[:40]); return ok(live.snapshot())
    except Exception as e:return fail(e)
@app.get('/api/stocks')
def stocks():
    try:
        u=request.args.get('universe','nifty50'); p=request.args.get('period','6m'); top=_int('top',20,1,100); mg=_float('min_growth')
        df,b=core.screen_stocks(u,p,mg,top,request.args.get('consistent')=='1',False,False)
        return ok(rec(df),buckets=b)
    except Exception as e:return fail(e,500)
@app.get('/api/funds')
def funds():
    try:
        q=(request.args.get('q') or 'flexi cap').strip(); p=request.args.get('period','1y'); top=_int('top',20,1,50); scan=_int('scan',40,5,200); mg=_float('min_growth')
        df,b=core.screen_funds(q,p,mg,top,scan,False,False); return ok(rec(df),buckets=b)
    except Exception as e:return fail(e,500)
@app.get('/api/suggest/equity')
def suggest_eq():
    try:return ok(core.suggest_equities((request.args.get('q') or '').strip(),12))
    except Exception as e:return fail(e)
@app.get('/api/suggest/fund')
def suggest_mf():
    try:return ok(core.suggest_funds((request.args.get('q') or '').strip(),12))
    except Exception as e:return fail(e)
@app.get('/api/analyze/equity')
def analyze_eq():
    try:
        t=(request.args.get('ticker') or '').strip().upper()
        if not t:return fail('Stock name/symbol enter karein.')
        if not t.startswith('^') and '.' not in t:t+='.NS'
        r=core.analyze_equity(t); s=r['series'].dropna()
        return ok({'kind':'equity','name':r['name'],'id':r['id'],'row':r['row'],'info':r.get('info',{}),'news':r.get('news',[])[:12],'dividends':r.get('dividends',[]),'chart':[{'date':d.strftime('%Y-%m-%d'),'value':round(float(v),4)} for d,v in s.tail(1600).items()]})
    except Exception as e:return fail(e)
@app.get('/api/analyze/fund')
def analyze_mf():
    try:
        c=(request.args.get('code') or '').strip(); r=core.analyze_fund(c); s=r['series'].dropna()
        return ok({'kind':'fund','name':r['name'],'id':str(r['id']),'row':r['row'],'info':r.get('info',{}),'news':r.get('news',[])[:12],'chart':[{'date':d.strftime('%Y-%m-%d'),'value':round(float(v),4)} for d,v in s.tail(2000).items()]})
    except Exception as e:return fail(e)
@app.post('/api/export')
def export():
    try:
        j=request.get_json(silent=True) or {}; rows=j.get('rows') or []; title=(j.get('title') or 'Market Report')[:31]
        if not rows:return fail('Export ke liye data nahi hai.')
        bio=io.BytesIO()
        with pd.ExcelWriter(bio,engine='openpyxl') as w:
            pd.DataFrame(rows).to_excel(w,index=False,sheet_name=title)
            ws=w.book[title]; ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
            fill=PatternFill('solid',fgColor='17365D'); white='FFFFFF'; thin=Side(style='thin',color='D9E2F3')
            for c in ws[1]: c.font=Font(bold=True,color=white); c.fill=fill; c.alignment=Alignment(horizontal='center')
            for row in ws.iter_rows():
                for c in row: c.border=Border(bottom=thin); c.alignment=Alignment(vertical='center')
            for i,col in enumerate(ws.columns,1):
                width=min(42,max(11,max(len(str(c.value or '')) for c in col)+2)); ws.column_dimensions[get_column_letter(i)].width=width
        bio.seek(0); return send_file(bio,as_attachment=True,download_name='Market_Screener_Report.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    except Exception as e:return fail(e)
if __name__=='__main__': app.run(host='0.0.0.0',port=int(os.getenv('PORT','5000')),threaded=True)
